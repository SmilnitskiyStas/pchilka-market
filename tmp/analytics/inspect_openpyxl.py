from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from io import BytesIO
from openpyxl import load_workbook

base = Path(r"C:\Users\loyalty\source\PchilkaWebApp\docs\Analytics")
fixed_dir = Path(r"C:\Users\loyalty\source\PchilkaWebApp\tmp\analytics\normalized")
fixed_dir.mkdir(parents=True, exist_ok=True)

for file in sorted(base.glob("*.xlsx")):
    print(f"\n### {file.name} size={file.stat().st_size}")
    with ZipFile(file) as z:
        ct = z.read("[Content_Types].xml").decode("utf-8", "replace")
        print("sharedStrings correct type:", "spreadsheetml.sharedStrings+xml" in ct)
        print("pivot parts:", [n for n in z.namelist() if "pivot" in n.lower()])
    fixed = fixed_dir / file.name
    with ZipFile(file) as src, ZipFile(fixed, "w", ZIP_DEFLATED) as dst:
        for item in src.infolist():
            name = "xl/sharedStrings.xml" if item.filename == "xl/SharedStrings.xml" else item.filename
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                text = data.decode("utf-8", "replace")
                text = text.replace('PartName="/xl/SharedStrings.xml"', 'PartName="/xl/sharedStrings.xml"')
                text = text.replace('ContentType="application/xml"', 'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"') if "sharedStrings.xml" in text and "spreadsheetml.sharedStrings+xml" not in text else text
                data = text.encode("utf-8")
            dst.writestr(name, data)
    wb = load_workbook(fixed, data_only=False, read_only=False)
    for ws in wb.worksheets:
        print(f"-- {ws.title}: {ws.max_row}x{ws.max_column}; merges={list(ws.merged_cells.ranges)}; tables={list(ws.tables.keys())}; pivots={len(getattr(ws, '_pivots', []))}")
        rows = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 25), min_col=1, max_col=min(ws.max_column, 30)):
            vals = [c.value for c in row]
            if any(v is not None for v in vals):
                while vals and vals[-1] is None:
                    vals.pop()
                rows.append(vals)
        for row in rows:
            print(repr(row))
