from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from collections import Counter
from openpyxl import load_workbook

root = Path(r"C:\Users\loyalty\source\PchilkaWebApp")
base = root / "docs" / "Analytics"
fixed_dir = root / "tmp" / "analytics" / "new_normalized"
fixed_dir.mkdir(parents=True, exist_ok=True)

def normalize_xlsx(source: Path, target: Path):
    with ZipFile(source) as src, ZipFile(target, "w", ZIP_DEFLATED) as dst:
        names = src.namelist()
        has_upper = "xl/SharedStrings.xml" in names
        for item in src.infolist():
            name = "xl/sharedStrings.xml" if item.filename == "xl/SharedStrings.xml" else item.filename
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                text = data.decode("utf-8", "replace")
                text = text.replace('PartName="/xl/SharedStrings.xml"', 'PartName="/xl/sharedStrings.xml"')
                if 'PartName="/xl/sharedStrings.xml" ContentType="application/xml"' in text:
                    text = text.replace(
                        'PartName="/xl/sharedStrings.xml" ContentType="application/xml"',
                        'PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"',
                    )
                data = text.encode("utf-8")
            dst.writestr(name, data)
    return has_upper

for file in sorted(base.glob("*.xlsx")):
    fixed = fixed_dir / file.name
    upper = normalize_xlsx(file, fixed)
    wb = load_workbook(fixed, data_only=False, keep_links=False)
    print(f"\n### {file.name} uppercaseSharedStrings={upper} sheets={wb.sheetnames}")
    for ws in wb.worksheets:
        formulas = []
        errors = []
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f": formulas.append((c.coordinate, c.value))
                if c.data_type == "e": errors.append((c.coordinate, c.value))
        print(f"sheet={ws.title} dimensions={ws.max_row}x{ws.max_column} merges={len(ws.merged_cells.ranges)} formulas={len(formulas)} errors={len(errors)}")
        print("top rows:")
        for r in range(1, min(ws.max_row, 7) + 1):
            vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column, 30) + 1)]
            while vals and vals[-1] is None: vals.pop()
            print(r, repr(vals))
        print("formula sample:", formulas[:20])
        print("bottom rows:")
        for r in range(max(1, ws.max_row - 3), ws.max_row + 1):
            vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column, 30) + 1)]
            while vals and vals[-1] is None: vals.pop()
            print(r, repr(vals))
        if ws.max_row >= 6:
            labels = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value]
            dupes = [k for k, v in Counter(labels).items() if v > 1]
            print(f"labels={len(labels)} duplicateLabels={dupes}")
