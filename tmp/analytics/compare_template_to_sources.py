from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

root = Path(r"C:\Users\loyalty\source\PchilkaWebApp")
base = root / "tmp" / "analytics" / "new_normalized"
old = load_workbook(base / "Старий з 25.08 по 31.08.26.xlsx", data_only=True)["TDSheet"]
new = load_workbook(base / "Новий з 01.09 по 07.09.26.xlsx", data_only=True)["TDSheet"]
sample = load_workbook(base / "Шок ціна з 01.09 по 08.09.26.xlsx", data_only=True)["TDSheet"]

def error_counts(ws):
    return Counter(c.value for row in ws.iter_rows() for c in row if c.data_type == "e")

print("old_errors", dict(error_counts(old)))
print("new_errors", dict(error_counts(new)))

old_rows = {old.cell(r, 1).value: r for r in range(6, old.max_row) if str(old.cell(r, 1).value or "").startswith("Магазин")}
new_rows = {new.cell(r, 1).value: r for r in range(6, new.max_row) if str(new.cell(r, 1).value or "").startswith("Магазин")}
sample_rows = {sample.cell(r, 1).value: r for r in range(4, sample.max_row) if str(sample.cell(r, 1).value or "").startswith("Магазин")}

def no_zero(value):
    return str(value).replace("0", "") if value is not None else None

matched_by_zero_stripping = 0
unmatched_sample_stores = []
for s in sample_rows:
    if s in old_rows:
        continue
    candidates = [o for o in old_rows if no_zero(o) == s]
    if len(candidates) == 1:
        matched_by_zero_stripping += 1
    else:
        unmatched_sample_stores.append(s)
print("sample_store_count", len(sample_rows))
print("exact_store_names", sum(1 for s in sample_rows if s in old_rows))
print("store_names_matching_after_zero_removal", matched_by_zero_stripping)
print("unmatched_sample_stores", unmatched_sample_stores)

header_pairs = []
for g in range(13):
    old_col = 2 + 3 * g
    sample_col = 2 + 4 * g
    header_pairs.append((old.cell(1, old_col).value, sample.cell(1, sample_col).value, old.cell(3, old_col).value, sample.cell(2, sample_col).value))
print("header_pairs")
for pair in header_pairs:
    print(pair)

# Compare every sample body cell to expected source values, matching damaged sample labels after zero removal.
field_mismatches = Counter()
zero_removal_matches = Counter()
examples = []
for sample_store, sample_row in sample_rows.items():
    source_store = sample_store if sample_store in old_rows else next((o for o in old_rows if no_zero(o) == sample_store), None)
    if source_store is None or source_store not in new_rows:
        continue
    orow, nrow = old_rows[source_store], new_rows[source_store]
    for g in range(13):
        scol = 2 + 4 * g
        ocol = 2 + 3 * g
        ncol = 2 + 2 * g
        old_qty = old.cell(orow, ocol + 1).value
        new_qty = new.cell(nrow, ncol + 1).value
        balance = new.cell(nrow, ncol).value
        coeff = (new_qty / old_qty) if isinstance(old_qty, (int, float)) and old_qty != 0 and isinstance(new_qty, (int, float)) else None
        expected = [old_qty, new_qty, coeff, balance]
        for f, exp in enumerate(expected):
            actual = sample.cell(sample_row, scol + f).value
            if actual != exp:
                field = ["old_qty", "new_qty", "coefficient", "new_balance"][f]
                field_mismatches[field] += 1
                if no_zero(exp) == no_zero(actual):
                    zero_removal_matches[field] += 1
                if len(examples) < 15:
                    examples.append((sample_store, g + 1, field, actual, exp))
print("field_mismatches", dict(field_mismatches))
print("mismatches_explained_by_zero_removal", dict(zero_removal_matches))
print("mismatch_examples")
for x in examples:
    print(x)
