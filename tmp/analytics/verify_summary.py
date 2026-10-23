from pathlib import Path
from openpyxl import load_workbook

root = Path(r"C:\Users\loyalty\source\PchilkaWebApp")
src = root / "tmp" / "analytics" / "normalized"
final_path = root / "outputs" / "analytics-shock-price-20260908" / "Шок ціна з 01.09 по 07.09.26 - зведення.xlsx"

old_ws = load_workbook(src / "Отчет-25-31,08.xlsx", data_only=True)["TDSheet"]
new_ws = load_workbook(src / "Отчет-1-7,09.xlsx", data_only=True)["TDSheet"]
final_ws = load_workbook(final_path, data_only=False)["TDSheet"]

stores = [old_ws.cell(r, 1).value for r in range(6, old_ws.max_row) if old_ws.cell(r, 1).value]
old_by_store = {old_ws.cell(r, 1).value: r for r in range(6, old_ws.max_row)}
new_by_store = {new_ws.cell(r, 1).value: r for r in range(6, new_ws.max_row)}

mismatches = []
checks = 0
for out_row, store in enumerate(stores, start=4):
    if final_ws.cell(out_row, 1).value != store:
        mismatches.append((out_row, 1, final_ws.cell(out_row, 1).value, store))
    for group in range(14):
        src_start = 2 + group * 3
        out_start = 2 + group * 5
        old_row = old_by_store[store]
        new_row = new_by_store.get(store)
        expected = [
            old_ws.cell(old_row, src_start).value,
            old_ws.cell(old_row, src_start + 1).value,
            new_ws.cell(new_row, src_start + 1).value if new_row else None,
            old_ws.cell(old_row, src_start + 2).value,
        ]
        actual = [
            final_ws.cell(out_row, out_start).value,
            final_ws.cell(out_row, out_start + 1).value,
            final_ws.cell(out_row, out_start + 2).value,
            final_ws.cell(out_row, out_start + 4).value,
        ]
        checks += 4
        for pos, (a, e) in enumerate(zip(actual, expected)):
            if a != e:
                mismatches.append((out_row, out_start + [0, 1, 2, 4][pos], a, e))
        formula = final_ws.cell(out_row, out_start + 3).value
        if not (isinstance(formula, str) and formula.startswith("=IF(OR(")):
            mismatches.append((out_row, out_start + 3, formula, "comparison formula"))

assert final_ws.max_row == 42, final_ws.max_row
assert final_ws.max_column == 71, final_ws.max_column
assert final_ws.freeze_panes == "B4", final_ws.freeze_panes
print(f"checked_static_cells={checks}")
print(f"stores={len(stores)} groups=14 dimensions={final_ws.max_row}x{final_ws.max_column}")
print(f"mismatches={len(mismatches)}")
if mismatches:
    print(mismatches[:20])
    raise SystemExit(1)
